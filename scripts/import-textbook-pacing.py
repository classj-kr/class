"""Extract acquired publisher workbooks into source-traceable curriculum data.

Read-only XLSX extraction; never executes workbook formulas or macros.
Usage: python scripts/import-textbook-pacing.py D:/Downloads
"""
import hashlib
import json
import re
import shutil
import sys
import subprocess
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'game-hub-server/data/textbooks'
RAW = ROOT / 'references/textbooks/raw'
CATALOG = json.loads((OUT / 'catalog-2022.json').read_text(encoding='utf-8'))
VISANG = {
 '수학 1-1': '432955', '수학 1-2': '444475', '수학 2-1': '432961', '수학 2-2': '444482',
 '수학(방) 3-1': '458914', '수학(방) 3-2': '521152', '수학(방) 4-1': '458941', '수학(방) 4-2': '521181', '수학(방) 5-1': '560350', '수학(방) 6-1': '560386',
 '사회(설) 3-1': '458541', '사회(설) 3-2': '458542', '사회(설) 4-1': '458579', '사회(설) 4-2': '458580', '사회(설) 5-1': '563321', '사회(설) 6-1': '563358',
 '과학(강) 3-1': '458034', '과학(강) 3-2': '458035', '과학(강) 4-1': '458063', '과학(강) 4-2': '458064', '과학(강) 5-1': '563647', '과학(강) 6-1': '563678',
 '음악(주) 3-1': '458882', '음악(주) 3-2': '490050', '음악(주) 4-1': '458899', '음악(주) 4-2': '490051', '음악(주) 5': '564469', '음악(주) 6': '564487',
 '미술(이) 3': '458634', '미술(이) 4': '458652', '미술(이) 5': '564280', '미술(이) 6': '564253',
 '체육(송) 3': '458434', '체육(송) 4': '458419', '체육(송) 5': '564502', '체육(송) 6': '564525',
 '영어(우) 5': '564872', '영어(우) 6': '564837', '실과(송) 5-2': '564366', '실과(송) 6-1': '564381',
}
AUTHORS = {
 '비상교육': {'수학':'방정숙','사회':'설규주','과학':'강석진','영어':'우길주','음악':'주대창','미술':'이재영','체육':'송지환','실과':'송현순'},
 '아이스크림미디어': {'수학':'김성여','사회':'한춘희','과학':'박일우','영어':'박유미','음악':'조순이','미술':'손지현','체육':'김명숙','실과':'정영식'},
}

def text(value):
    return '' if value is None else str(value).strip()

def periods(value):
    raw = text(value).split('/')[0].replace('～','~')
    match = re.fullmatch(r'(\d+)\s*[-~]\s*(\d+)', raw)
    if match:
        start, end = map(int,match.groups())
        return list(range(start,end+1)) if 0 < start <= end < 200 else None
    return [int(raw)] if raw.isdigit() and int(raw)>0 else None

def metadata(p):
    if '비상교육' in p.name:
        found = re.search(r'초등_(.+?)_진도표',p.name)
        if not found or found[1] not in VISANG:
            return None
        label = found[1]
        m = re.fullmatch(r'([가-힣]+)(?:\([가-힣]+\))? ([1-6])(?:-([12]))?',label)
        subject, grade, semester = m.groups()
        cid = VISANG[label]
        if subject=='미술' and grade in ('3','4') and '나이스' not in p.name:
            cid = {'3':'459757','4':'459759'}[grade]
        return {'publisher':'비상교육','subject':subject,'grade':int(grade),'semester':int(semester) if semester else None,
                'sourceId':cid,'sourceUrl':'https://e.vivasam.com/class/management/materials/list','sourceLabel':label,
                'curriculumVerification':'publisher-textbook-comparison-pending'}
    m = re.fullmatch(r'\(([가-힣]+)\)([3-6])학년 ([12])학기_진도표\(나이스 업로드용\)\.xlsx',p.name)
    if not m:
        return None
    subject, grade, semester = m.groups()
    grade, semester = int(grade),int(semester)
    page = (26579 if semester==1 else 26580) if grade<5 else ((26783 if semester==1 else 26784) if subject in ('미술','체육','실과') else (26790 if semester==1 else 26791))
    return {'publisher':'아이스크림미디어','subject':subject,'grade':grade,'semester':semester,
            'sourceId':str(page),'sourceUrl':f'https://text.i-scream.co.kr/user/information/communityView.do?no={page}',
            'sourceLabel':f'{2025 if grade<5 else 2026}학년도 {semester}학기 공식 진도표','curriculumVerification':'publisher-2025-2026-edition'}

def run(directory):
    plans, sources, issues = [], [], []
    RAW.mkdir(parents=True,exist_ok=True)
    for p in sorted(Path(directory).glob('*진도표*.xlsx')):
        meta = metadata(p)
        if not meta:
            continue
        digest = hashlib.sha256(p.read_bytes()).hexdigest()
        source_key = ('vivasam-' if meta['publisher']=='비상교육' else 'iscream-') + meta['sourceId'] + '-' + digest[:10]
        shutil.copy2(p,RAW / (source_key+'.xlsx'))
        matches = [e for e in CATALOG['editions'] if meta['grade'] in e['grades'] and e['subject']==meta['subject'] and ((meta['grade']<3 and e['approvalType']=='국정') or (e['publisher']==meta['publisher'] and e['leadAuthor']==AUTHORS[meta['publisher']][meta['subject']]))]
        if len(matches)!=1:
            raise ValueError(f'Unresolved edition: {p.name}: {len(matches)}')
        edition = matches[0]
        workbook = openpyxl.load_workbook(p,data_only=True)
        file_plans = []
        for sheet in workbook:
            # Expand merged cells only within their actual merge ranges.
            merged = {}
            for area in sheet.merged_cells.ranges:
                value = sheet.cell(area.min_row,area.min_col).value
                for row in range(area.min_row,area.max_row+1):
                    for col in range(area.min_col,area.max_col+1):
                        merged[(row,col)] = value
            headers = [text(c.value).replace(' ','') for c in sheet[1]]
            if sheet.max_row <= 1 and not any(headers):
                continue
            neis = '학습내용' in headers
            if not neis and '학습주제' not in headers:
                issues.append({'sourceId':source_key,'sheet':sheet.title,'issue':'unrecognized-header'})
                continue
            lessons = []
            for row_no in range(2,sheet.max_row+1):
                values = [merged.get((row_no,col),sheet.cell(row_no,col).value) for col in range(1,sheet.max_column+1)]
                cells = dict(zip(headers,values))
                topic = text(cells.get('학습내용' if neis else '학습주제'))
                if not topic:
                    continue
                unit = text(cells.get('단원' if neis else '단원명(소주제)'))
                period_text = text(cells.get('해당차시' if neis else '차시'))
                semester_text = text(cells.get('학기')) if neis else ''
                semester = int(semester_text) if semester_text in ('1','2') else meta['semester']
                row_grade = text(cells.get('학년'))
                if neis and row_grade and row_grade!=str(meta['grade']):
                    issues.append({'sourceId':source_key,'sheet':sheet.title,'row':row_no,'issue':'grade-mismatch','value':row_grade})
                standards = text(cells.get('(핵심)성취기준'))
                numbers = periods(period_text)
                lessons.append({'id':f'{source_key}-{sheet.title}-{row_no}','sourceRow':row_no,'sequence':len(lessons)+1,
                    'semester':semester,'unit':unit,'domain':text(cells.get('대단원(대주제)')) or None,'topic':topic,
                    'periodText':period_text,'periodNumbers':numbers,'suggestedPeriods':len(numbers) if numbers else None,
                    'unitPeriodsText':text(cells.get('전체차시')) if neis else (period_text.split('/')[1] if '/' in period_text else ''),
                    'pages':text(cells.get('쪽수' if neis else '교과서쪽')) or None,'supplementaryPages':text(cells.get('보조쪽수')) or None,
                    'standardCodes':re.findall(r'\[([^\]]+)\]',standards),
                    'activities':text(cells.get('학습활동내용')) or None,'materials':text(cells.get('준비물')) or None})
            if lessons:
                plan = {'id':source_key+'-'+str(len(file_plans)+1),'editionId':edition['id'],'curriculum':'2022',**meta,
                    'sourceFile':p.name,'sha256':digest,'sheet':sheet.title,'format':'neis' if neis else 'teaching-plan',
                    'lessons':lessons}
                plans.append(plan)
                file_plans.append(plan['id'])
        sources.append({**meta,'id':source_key,'filename':p.name,'sha256':digest,'bytes':p.stat().st_size,'retrievedOn':'2026-10-08','plans':file_plans})
    OUT.mkdir(parents=True,exist_ok=True)
    for publisher, short in [('비상교육','visang'),('아이스크림미디어','iscream')]:
        (OUT/f'pacing-{short}-2022.json').write_text(json.dumps({'schemaVersion':1,'publisher':publisher,'plans':[p for p in plans if p['publisher']==publisher]},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (ROOT/'references/textbooks/acquired-sources.json').write_text(json.dumps({'sources':sources,'issues':issues},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    subprocess.run(['node', str(ROOT/'scripts/build-textbook-coverage.cjs')], check=True)
    print(json.dumps({'files':len(sources),'plans':len(plans),'lessonRows':sum(len(p['lessons']) for p in plans),'issues':issues},ensure_ascii=False))

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    run(sys.argv[1])
