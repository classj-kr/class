"""Cache text from verified local originals; preserve extraction limitations."""
import argparse,hashlib,importlib.util,json,re,zipfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];REF=ROOT/'references/textbooks'
OUT=ROOT/'tmp/textbook-research/tselpa-assessment-text';OUT.mkdir(parents=True,exist_ok=True)
spec=importlib.util.spec_from_file_location('hwp_text',ROOT/'scripts/extract-assessment-text.py')
helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)

def read(p):return json.loads(p.read_text(encoding='utf-8'))
def write(p,d):p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def clean(t):return re.sub(r'\s+',' ',t.replace('\x00',' ')).strip()

def run():
    parser=argparse.ArgumentParser();parser.add_argument('--provider',choices=['tselpa','jihak'],default='tselpa');provider=parser.parse_args().provider
    out=ROOT/f'tmp/textbook-research/{provider}-assessment-text';out.mkdir(parents=True,exist_ok=True)
    reports=[]
    for item in read(REF/f'assessment-index-{provider}-2022.json')['assessments']:
        p=REF/'raw'/item['sourceFile'];cache=out/(item['id']+'.json')
        assert hashlib.sha256(p.read_bytes()).hexdigest()==item['sha256'],item['id']
        if cache.exists() and read(cache).get('sha256')==item['sha256'] and read(cache).get('method')!='unsupported-format':
            result=read(cache)
        else:
            result=dict(id=item['id'],sha256=item['sha256'],sourceFile=item['sourceFile'])
            try:
                if p.suffix.lower() in ['.hwp','.hwpx']:
                    result.update(method='hwp-paragraph-text',paragraphs=[clean(t) for t in helper.paragraphs(p) if clean(t)],
                                  limitations=['Equations, images, and table geometry are not represented.'])
                elif p.suffix.lower()=='.pdf':
                    from pypdf import PdfReader
                    pages=[p.extract_text() or '' for p in PdfReader(p).pages]
                    result.update(method='pdf-page-text',pages=pages,limitations=['Reading order and visual content require review.'])
                elif p.suffix.lower()=='.zip':
                    with zipfile.ZipFile(p) as z:
                        result.update(method='zip-member-inventory',members=[dict(name=n.filename,bytes=n.file_size) for n in z.infolist() if not n.is_dir()])
                elif p.suffix.lower()=='.xlsx':
                    from openpyxl import load_workbook
                    workbook=load_workbook(p,read_only=True,data_only=True)
                    result.update(method='xlsx-cell-values',sheets=[dict(name=s.title,rows=[[v if v is None or isinstance(v,(str,int,float,bool)) else str(v) for v in row] for row in s.iter_rows(values_only=True)]) for s in workbook.worksheets],
                                  limitations=['Formula expressions and cell formatting are not represented.'])
                    workbook.close()
                elif p.suffix.lower()=='.mp3':
                    result.update(method='audio-original-only',limitations=['Audio has been preserved; no transcript or listening review has been produced.'])
                else:result.update(method='unsupported-format')
                result['status']='text-cached-unreviewed' if any(k in result for k in ['paragraphs','pages','sheets']) else 'archive-inventoried' if 'members' in result else 'audio-preserved-unreviewed' if result['method']=='audio-original-only' else 'unsupported-format'
            except Exception as exc:
                result.update(status='extraction-failed',errorType=type(exc).__name__)
            write(cache,result)
        reports.append(dict(id=item['id'],sourceFile=item['sourceFile'],sha256=item['sha256'],status=result['status'],method=result.get('method'),
                            paragraphCount=len(result.get('paragraphs',[])),pageCount=len(result.get('pages',[])),memberCount=len(result.get('members',[]))))
    write(REF/f'assessment-{provider}-text-review.json',dict(schemaVersion=1,documents=reports,reviewStatus='text-extraction-only-not-structured-rubrics',
        cacheDirectory=f'tmp/textbook-research/{provider}-assessment-text',limitations='Text extraction does not establish visual accuracy or permission to republish.'))
    from collections import Counter
    print(json.dumps(dict(documents=len(reports),statuses=dict(Counter(r['status'] for r in reports))),ensure_ascii=False))
if __name__=='__main__':run()
