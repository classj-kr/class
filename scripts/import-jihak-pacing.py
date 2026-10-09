"""Read the five authenticated official Tsolution pacing ZIP downloads."""
import hashlib
import io
import json
import re
import shutil
import subprocess
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'game-hub-server/data/textbooks'
RAW=ROOT/'references/textbooks/raw'
CATALOG=json.loads((OUT/'catalog-2022.json').read_text(encoding='utf-8'))
AUTHORS={'수학':'강문봉','사회':'이간용','과학':'권치순','음악':'허정미','미술':'송미영','체육':'신기철','실과':'김차명'}
URL='https://tsol.jihak.co.kr/classResourceList.ez?leftMenuSeq=RESOURCE_PROGRESS&leftSubMenuSeq=a0b2a226ce3d48008a7a514dd314ad77'

def text(v):return '' if v is None else str(v).strip()

def periods(v):
    m=re.fullmatch(r'(\d+)(?:\s*[-~～]\s*(\d+))?',text(v))
    if not m:return None
    a,b=int(m[1]),int(m[2] or m[1]);return list(range(a,b+1)) if 0<a<=b<300 else None

def run(directory):
    plans,sources=[],[]
    for file in sorted(Path(directory).glob('연간지도계획*zip')):
        sha=hashlib.sha256(file.read_bytes()).hexdigest()
        shutil.copy2(file,RAW/f'jihak-batch-{sha[:12]}.zip')
        for name in (archive:=zipfile.ZipFile(file)).namelist():
            payload=archive.read(name);digest=hashlib.sha256(payload).hexdigest();sid=f'jihak-{digest[:16]}'
            source={'id':sid,'publisher':'지학사','sourceUrl':URL,'archiveFile':file.name,'archiveSha256':sha,
                    'filename':name,'sha256':digest,'retrievedOn':'2026-10-09','plans':[]}
            if not name.endswith('.xlsx'):
                source['status']='supplementary-science-tracking-form';sources.append(source);continue
            for sheet in openpyxl.load_workbook(io.BytesIO(payload),data_only=True):
                headers=[re.sub(r'\s+','',text(c.value)) for c in sheet[1]]
                topic_key='학습내용' if '학습내용' in headers else '소단원'
                assert topic_key in headers
                merged={}
                for a in sheet.merged_cells.ranges:
                    for r in range(a.min_row,a.max_row+1):
                        for c in range(a.min_col,a.max_col+1):merged[r,c]=sheet.cell(a.min_row,a.min_col).value
                groups=defaultdict(list)
                for rn in range(2,sheet.max_row+1):
                    cells=dict(zip(headers,[merged.get((rn,c),sheet.cell(rn,c).value) for c in range(1,len(headers)+1)]))
                    topic=text(cells.get(topic_key))
                    if not topic:continue
                    grade=int(cells['학년']);semester=int(cells['학기']) or None;subject=text(cells['편제'])
                    source_subject=subject
                    if name.startswith('(통합교과_') and subject=='통합':subject='통합교과'
                    assert grade in range(1,7) and semester in (None,1,2)
                    numbers=periods(cells.get('해당차시'))
                    groups[grade,semester,subject].append({'id':f'{sid}-{sheet.title}-{rn}','sourceRow':rn,
                        'sequence':len(groups[grade,semester,subject])+1,'semester':semester,'unit':text(cells.get('단원')),
                        'sourceSubject':source_subject,'isDiscretionary':'담임재량' in text(cells.get('단원')),
                        'subunit':text(next((cells[k] for k in ('중단원','소단원','제재명','주제','수업묶음') if cells.get(k)),None)) or None,
                        'topic':topic,'periodText':text(cells.get('해당차시')),'periodNumbers':numbers,
                        'suggestedPeriods':len(numbers) if numbers else None,'unitPeriodsText':text(cells.get('전체차시')),
                        'pages':text(cells.get('쪽수')) or None,'supplementaryPages':text(next((v for k,v in cells.items() if k.startswith('보조쪽수')),None)) or None,
                        'standardCodes':[],'activities':None,'materials':text(cells.get('준비물')) or None,
                        'granularity':'subunit' if topic_key=='소단원' else 'lesson'})
                for (grade,semester,subject),lessons in groups.items():
                    national=subject in ('국어','통합교과')
                    matches=[e for e in CATALOG['editions'] if grade in e['grades'] and e['subject']==subject and
                             ((national and e['approvalType']=='국정') or (not national and e['publisher']=='지학사' and e['leadAuthor']==AUTHORS[subject]))]
                    assert len(matches)==1,(name,grade,subject)
                    # Integrated units directly match the 2022 national volume titles.
                    if subject=='통합교과':
                        units={l['unit'] for l in lessons if not l['isDiscretionary']}
                        volume_titles=[v['title'] for v in matches[0]['volumes'] if v['grade']==grade]
                        assert all(any(u in title for title in volume_titles) for u in units),(name,units)
                    pid=f'{sid}-{grade}-{semester or 0}'
                    plans.append({'id':pid,'editionId':matches[0]['id'],'curriculum':'2022','publisher':'지학사','grade':grade,'semester':semester,
                                  'subject':subject,'sourceId':sid,'sourceUrl':URL,'sourceFile':name,'sourceArchive':file.name,
                                  'sha256':digest,'sheet':sheet.title,'format':'neis','curriculumVerification':'publisher-2022-label-and-national-catalog',
                                  'lessons':lessons})
                    source['plans'].append(pid)
            source['status']='normalized';sources.append(source)
    (OUT/'pacing-jihak-2022.json').write_text(json.dumps({'schemaVersion':1,'publisher':'지학사','plans':plans},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (ROOT/'references/textbooks/jihak-acquired-sources.json').write_text(json.dumps({'sources':sources},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    subprocess.run(['node',str(ROOT/'scripts/build-textbook-coverage.cjs')],check=True)
    print(json.dumps({'files':len(sources),'plans':len(plans),'lessonRows':sum(len(p['lessons']) for p in plans)}))

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8');run(sys.argv[1])
