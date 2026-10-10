"""Normalize reviewed official Kumsung annual plans, preserving allocation spans."""
import hashlib
import io
import json
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
import openpyxl
import pdfplumber

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'references/textbooks/수집작업/python-packages'))
from hwp5.xmlmodel import Hwp5File
OUT=ROOT/'game-hub-server/data/textbooks'
RAW=ROOT/'references/textbooks/raw'
CATALOG=json.loads((OUT/'catalog-2022.json').read_text(encoding='utf-8'))

def text(value):
    return re.sub(r'\s+',' ',str(value)).strip() if value is not None else ''

def xlsx_grid(sheet):
    grid={(r,c):text(sheet.cell(r+1,c+1).value) for r in range(sheet.max_row) for c in range(sheet.max_column)}
    anchors={rc:rc for rc in grid}
    for area in sheet.merged_cells.ranges:
        anchor=(area.min_row-1,area.min_col-1)
        for r in range(area.min_row-1,area.max_row):
            for c in range(area.min_col-1,area.max_col):grid[r,c]=grid[anchor];anchors[r,c]=anchor
    return grid,anchors,sheet.max_row

def hwp_grid(file):
    buffer=io.BytesIO();doc=Hwp5File(str(file))
    try:doc.xmlevents().dump(buffer)
    finally:doc.close()
    tables=ET.fromstring(buffer.getvalue()).findall('.//TableBody')
    table=max(tables,key=lambda t:int(t.get('rows')))
    grid,anchors={},{}
    for cell in table.findall('./TableRow/TableCell'):
        r,c,rs,cs=[int(cell.get(k)) for k in ('row','col','rowspan','colspan')]
        value=text(' '.join(''.join(p.itertext()) for p in cell.findall('./Paragraph')))
        for rr in range(r,r+rs):
            for cc in range(c,c+cs):grid[rr,cc]=value;anchors[rr,cc]=(r,c)
    return grid,anchors,int(table.get('rows'))

def run(directory):
    directory=Path(directory);plans=[];sources=[]
    specs=[('음악',5,'music','김용희','음악5_공통자료_연간지도계획_xlxs.xlsx')]
    specs += [('미술',g,'art','김정선',f'미술{g}_공통자료_연간지도계획.Zip') for g in (5,6)]
    specs += [('체육',g,'physicalEdu','이제행',f'체육{g}_공통자료_연간지도계획.hwp') for g in (5,6)]
    specs += [('실과',g,'practical','류청산',f'실과{g}_공통자료_연간지도계획.zip') for g in (5,6)]
    for subject,source_grade,slug,author,filename in specs:
        file=directory/filename;digest=hashlib.sha256(file.read_bytes()).hexdigest()
        sid=f'kumsung-{slug}-{source_grade}-{digest[:10]}'
        shutil.copy2(file,RAW/(sid+file.suffix.lower()))
        url=f'https://newtext2025.kumsung.co.kr/textbook/element/{slug}/{source_grade}grade/'
        source={'id':sid,'filename':filename,'sourceUrl':url,'publisher':'금성출판사','sha256':digest,'retrievedOn':'2026-10-09','plans':[]}
        if file.suffix.lower()=='.zip':
            with zipfile.ZipFile(file) as z:
                member=next(n for n in z.namelist() if n.endswith('.xlsx'));payload=io.BytesIO(z.read(member));source['archiveMember']=member
            sheets=[(source_grade,s.title,*xlsx_grid(s)) for s in openpyxl.load_workbook(payload,data_only=True)]
        elif file.suffix=='.xlsx':
            sheets=[(int(s.title[0]),s.title,*xlsx_grid(s)) for s in openpyxl.load_workbook(file,data_only=True)]
        else:sheets=[(source_grade,'HWP table',*hwp_grid(file))]
        for grade,sheet,grid,anchors,row_count in sheets:
            edition=next(e for e in CATALOG['editions'] if e['publisher']=='금성출판사' and e['subject']==subject and e['leadAuthor']==author and grade in e['grades'])
            pid=f'{sid}-{grade}'
            plan={'id':pid,'editionId':edition['id'],'curriculum':'2022','publisher':'금성출판사','subject':subject,'grade':grade,'semester':None,
                  'sourceId':sid,'sourceUrl':url,'sourceFile':filename,'sha256':digest,'sheet':sheet,'format':'teaching-plan',
                  'curriculumVerification':'publisher-2022-promotion','lessons':[]}
            for r in range(4 if subject=='실과' else 2 if subject=='미술' else 1,row_count):
                if subject=='음악':
                    unit,topic,period,pages,guide=[grid.get((r,c),'') for c in (0,1,2,4,5)];semester=None;pc=2;subunit=None
                    if unit=='합계' or not topic:continue
                    activities=grid.get((r,3))
                elif subject=='미술':
                    unit,subunit,topic,period,pages,guide=[grid.get((r,c),'') for c in (1,2,3,4,5,6)];pc=4
                    if unit=='총 차시' or not topic:continue
                    semester=int(grid[r,0][0]);activities=None
                elif subject=='체육':
                    unit,subunit,topic,period,pages,guide=[grid.get((r,c),'') for c in (2,3,4,5,6,7)];pc=5;semester=None;activities=None
                    if not topic or not period.isdigit():continue
                else:
                    unit,subunit,topic,period,pages,guide=[grid.get((r,c),'') for c in (3,4,5,2,7,8)];pc=2;semester=None;activities=None
                    if not topic or not period.isdigit():continue
                assert unit and period.isdigit(),(filename,r,unit,period)
                anchor=anchors[r,pc][0]
                allocation=int(period) if anchor==r else None
                lesson={'id':f'{pid}-r{r+1}','sourceRow':r+1,'sequence':len(plan['lessons'])+1,'semester':semester,'unit':unit,'subunit':subunit,
                        'topic':topic,'periodText':f'{period}차시 배정' if anchor==r else f'{period}차시 공동 배정',
                        'periodNumbers':None,'suggestedPeriods':allocation,'periodBasis':'allocation',
                        'pages':pages or None,'supplementaryPages':guide or None,'standardCodes':[],'activities':activities,'materials':None}
                if anchor!=r:lesson['sharedPeriodWith']=f'{pid}-r{anchor+1}'
                if subject in ('체육','실과'):lesson['sourceMonth']=grid.get((r,0))
                if subject=='실과':lesson['sourcePeriodText']=grid.get((r,6))
                plan['lessons'].append(lesson)
            total=sum(l['suggestedPeriods'] or 0 for l in plan['lessons'])
            assert total==(102 if subject=='체육' else 68),(filename,grade,total)
            plans.append(plan);source['plans'].append(pid)
        sources.append(source)
    file=RAW/'kumsung-art-3-4.pdf';digest=hashlib.sha256(file.read_bytes()).hexdigest();sid=f'kumsung-art-3-4-{digest[:10]}'
    url='https://newtext2025.kumsung.co.kr/textbook/element/art/3grade/'
    source={'id':sid,'filename':file.name,'sourceUrl':url,'publisher':'금성출판사','sha256':digest,'retrievedOn':'2026-10-09','plans':[]}
    with pdfplumber.open(file) as pdf:
        for page,grade in zip(pdf.pages,(3,4)):
            table_index,table=next((i,t) for i,t in enumerate(page.extract_tables(),1) if t[0][0]=='대단원명')
            edition=next(e for e in CATALOG['editions'] if e['publisher']=='금성출판사' and e['subject']=='미술' and grade in e['grades'])
            pid=f'{sid}-{grade}';unit=None;anchor_id=None;period=None
            plan={'id':pid,'editionId':edition['id'],'curriculum':'2022','publisher':'금성출판사','subject':'미술','grade':grade,'semester':None,
                  'sourceId':sid,'sourceUrl':url,'sourceFile':file.name,'sha256':digest,'format':'pdf-table',
                  'curriculumVerification':'publisher-2022-promotion','lessons':[]}
            for rn,row in enumerate(table[1:],2):
                if '총 차시' in text(row[0]) or not row[1]:continue
                if row[0]:unit=text(row[0])
                if row[3]:period=int(row[3]);anchor_id=f'{pid}-r{rn}'
                item={'id':f'{pid}-r{rn}','sourcePage':grade-2,'sourceTable':table_index,'sourceRow':rn,'sourceCells':row,
                      'sequence':len(plan['lessons'])+1,'semester':None,'unit':unit,'topic':text(row[1]),'activities':text(row[2]),
                      'granularity':'subunit','periodText':f'{period}차시 배정' if row[3] else f'{period}차시 공동 배정',
                      'periodNumbers':None,'suggestedPeriods':period if row[3] else None,'periodBasis':'allocation',
                      'pages':text(row[4]),'supplementaryPages':text(row[5]),'standardCodes':[],'materials':None}
                if not row[3]:item['sharedPeriodWith']=anchor_id
                plan['lessons'].append(item)
            assert sum(l['suggestedPeriods'] or 0 for l in plan['lessons'])==68
            plans.append(plan);source['plans'].append(pid)
    sources.append(source)
    # Both grade-specific promotion links currently deliver the exact same file.
    # Preserve it for textbook comparison instead of silently assigning it to both grades.
    for grade in (3,4):
        file=directory/f'공통_2022 개정 3~4학년 체육과 연간 지도 계획_{grade}학년.hwpx'
        digest=hashlib.sha256(file.read_bytes()).hexdigest();sid=f'kumsung-physicalEdu-{grade}-{digest[:10]}'
        shutil.copy2(file,RAW/(sid+'.hwpx'))
        sources.append({'id':sid,'filename':file.name,'publisher':'금성출판사','sha256':digest,'retrievedOn':'2026-10-09',
                        'sourceUrl':f'https://newtext2025.kumsung.co.kr/textbook/element/physicalEdu/{grade}grade/',
                        'plans':[],'status':'grade-comparison-pending','issue':'Grade 3 and 4 links provide byte-identical annual plans.'})
    (OUT/'pacing-kumsung-2022.json').write_text(json.dumps({'schemaVersion':1,'publisher':'금성출판사','plans':plans},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (ROOT/'references/textbooks/kumsung-acquired-sources.json').write_text(json.dumps({'sources':sources},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    subprocess.run(['node',str(ROOT/'scripts/build-textbook-coverage.cjs')],check=True)

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8');run(sys.argv[1])
