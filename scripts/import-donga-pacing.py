"""Import official Douclass downloads, taking grade/semester from worksheet rows."""
import hashlib
import json
import re
import shutil
import subprocess
import sys
from collections import defaultdict
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'game-hub-server/data/textbooks'
RAW = ROOT / 'references/textbooks/raw'
CATALOG = json.loads((OUT / 'catalog-2022.json').read_text(encoding='utf-8'))
AUTHORS = {'수학':'나귀수', '사회':'박영석', '과학':'장신호', '음악':'최미영', '미술':'김정희', '체육':'최의창'}
# IDs observed on the official 2022 curriculum resource page, not guessed URLs.
SOURCE_IDS = {
 ('수학',3,1):'99875', ('수학',3,2):'120750', ('수학',4,1):'103679', ('수학',4,2):'120763',
 ('수학',5,1):'124966', ('수학',5,2):'156591', ('수학',6,1):'125090', ('수학',6,2):'156187',
 ('사회',3,1):'105509', ('사회',4,1):'103368', ('사회',5,1):'126884', ('사회',5,2):'157584', ('사회',6,1):'126933', ('사회',6,2):'157922',
 ('과학',3,1):'103813', ('과학',3,2):'103353', ('과학',4,1):'67327', ('과학',4,2):'103962',
 ('영어',3,None):'68066', ('영어',4,None):'68079', ('영어',5,None):'128559', ('영어',6,None):'128561',
 ('음악',3,None):'92600', ('음악',4,None):'92601', ('음악',5,None):'125407', ('음악',6,None):'125425',
 ('미술',3,None):'64371', ('미술',6,None):'126432',
 ('체육',3,None):'63530', ('체육',4,None):'64816', ('체육',5,None):'124368', ('체육',6,None):'124694',
}

def clean(v):
    return '' if v is None else str(v).strip()

def period_numbers(value):
    raw = clean(value).replace('～', '~')
    m = re.fullmatch(r'(\d+)\s*[-~]\s*(\d+)', raw)
    if m:
        a, b = map(int, m.groups())
        return list(range(a,b+1)) if 0 < a <= b < 200 else None
    return [int(raw)] if raw.isdigit() and int(raw) > 0 else None

def run(directory):
    plans, sources = [], []
    RAW.mkdir(parents=True, exist_ok=True)
    for file in sorted(Path(directory).glob('*동아*.xlsx')):
        m = re.search(r'_(수학|사회|과학|영어|음악|미술|체육)(?:과)?\s*([3-6])(?:학년)?(?:\s*[- ]\s*([12])(?:학기)?)?', file.name)
        if not m:
            raise ValueError(f'Unknown filename: {file.name}')
        subject, file_grade, file_term = m.groups()
        key = (subject, int(file_grade), int(file_term) if file_term else None)
        cid = SOURCE_IDS[key]
        digest = hashlib.sha256(file.read_bytes()).hexdigest()
        source_key = f'donga-{cid}-{digest[:10]}'
        shutil.copy2(file, RAW / f'{source_key}.xlsx')
        source = {'id':source_key, 'sourceId':cid, 'sourceUrl':'https://ele.douclass.com/curriculum',
                  'publisher':'동아출판', 'filename':file.name, 'sha256':digest, 'bytes':file.stat().st_size,
                  'retrievedOn':'2026-10-09', 'curriculumVerification':'publisher-2022-filter', 'plans':[]}
        for sheet in openpyxl.load_workbook(file, data_only=True):
            headers = [re.sub(r'\s+', '', clean(c.value)) for c in sheet[1]]
            assert '학습내용' in headers, (file.name, sheet.title, headers)
            merged = {}
            for area in sheet.merged_cells.ranges:
                for row in range(area.min_row, area.max_row+1):
                    for col in range(area.min_col, area.max_col+1):
                        merged[row,col] = sheet.cell(area.min_row, area.min_col).value
            groups = defaultdict(list)
            sheet_grades = {int(r[headers.index('학년')]) for r in list(sheet.values)[1:] if r[headers.index('학년')] is not None}
            for rn in range(2, sheet.max_row+1):
                cells = dict(zip(headers, [merged.get((rn,c),sheet.cell(rn,c).value) for c in range(1,len(headers)+1)]))
                topic = clean(cells.get('학습내용'))
                if not topic:
                    continue
                # English 6 semester 2 row 27 omits grade; its sheet consistently says 6.
                missing_grade = cells['학년'] is None
                if missing_grade:
                    assert len(sheet_grades)==1, (file.name,sheet.title,rn)
                grade = next(iter(sheet_grades)) if missing_grade else int(cells['학년'])
                semester = int(cells['학기'])
                assert clean(cells['편제']) == subject and grade in range(3,7) and semester in (1,2), (file.name, rn)
                period = clean(cells.get('해당차시'))
                numbers = period_numbers(period)
                # Music 5/6 uses allocated hours per topic; music 3/4 uses numbered lessons.
                allocation = subject == '음악' and grade >= 5
                lesson = {'id':f'{source_key}-{sheet.title}-{rn}', 'sourceRow':rn,
                          'sequence':len(groups[grade,semester])+1, 'semester':semester,
                          'unit':clean(cells.get('단원')), 'topic':topic, 'domain':None,
                          'periodText':f'{period}차시 배정' if allocation else period,
                          'periodNumbers':None if allocation else numbers,
                          'suggestedPeriods':int(period) if allocation and period.isdigit() else (len(numbers) if numbers else None),
                          'periodBasis':'allocation' if allocation else 'source-numbering',
                          'unitPeriodsText':clean(cells.get('전체차시')),
                          'pages':clean(cells.get('쪽수')) or None,
                          'supplementaryPages':clean(cells.get('보조쪽수')) or None,
                          'standardCodes':[], 'activities':None, 'materials':clean(cells.get('준비물')) or None}
                if missing_grade:
                    lesson['sourceCorrections'] = ['missing-grade-inferred-from-unanimous-sheet-grade']
                groups[grade,semester].append(lesson)
            for (grade,semester), lessons in groups.items():
                author = ('윤여범' if grade<5 else '정은숙') if subject=='영어' else AUTHORS[subject]
                editions = [e for e in CATALOG['editions'] if e['publisher']=='동아출판' and e['subject']==subject and grade in e['grades'] and e['leadAuthor']==author]
                assert len(editions)==1, (subject,grade,author)
                pid = f'{source_key}-{sheet.title}-{grade}-{semester}'
                plans.append({'id':pid, 'editionId':editions[0]['id'], 'curriculum':'2022', 'publisher':'동아출판',
                              'subject':subject, 'grade':grade, 'semester':semester, 'sourceId':cid,
                              'sourceUrl':source['sourceUrl'], 'sourceFile':file.name, 'sha256':digest,
                              'curriculumVerification':source['curriculumVerification'], 'sheet':sheet.title,
                              'format':'neis', 'lessons':lessons})
                source['plans'].append(pid)
        sources.append(source)
    (OUT/'pacing-donga-2022.json').write_text(json.dumps({'schemaVersion':1,'publisher':'동아출판','plans':plans},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (ROOT/'references/textbooks/donga-acquired-sources.json').write_text(json.dumps({'sources':sources},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    subprocess.run(['node',str(ROOT/'scripts/build-textbook-coverage.cjs')],check=True)
    print(json.dumps({'files':len(sources),'plans':len(plans),'lessonRows':sum(len(p['lessons']) for p in plans)}))

if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    run(sys.argv[1])
